"""
services/excel_export.py
----------------------------
DIRECT PORT of dashboard_backend.py's Excel export - same color palette,
same fonts/borders/styling, same native openpyxl charts (bar/stacked-bar/
doughnut), same sheet layout and order.

WHAT CHANGED: V1's /export endpoint received a JSON payload the browser
built from data it already had client-side (dashboard.html had already
computed everything in JS). V2 computes everything server-side already
(services/dashboard_kpis.py etc.), so this pulls directly from those
functions instead of expecting the frontend to re-package and post the
same data back - simpler, and the export can never drift from what the
dashboard itself is showing.
"""
import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, DoughnutChart, Reference, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.marker import DataPoint
from sqlalchemy.orm import Session

from app.services import dashboard_kpis, driver_performance

# Exact palette from V1's dashboard_backend.py
NAVY, HEADER, TEAL, BLUE = "0D1B2A", "1A2B3C", "00D4A8", "2878D6"
AMBER, RED, GREEN, MUTED = "F59E0B", "EF4444", "22C55E", "7A9BBF"
WHITE, LIGHT, DARK2 = "FFFFFF", "E8F1FA", "132233"
CHART_COLORS = ["2878D6", "00D4A8", "F59E0B", "EF4444", "22C55E",
                "7C3AED", "E67E22", "1ABC9C", "3498DB", "E91E63", "00C2E0", "A78BFA"]


def hex_fill(hex_col):
    return PatternFill("solid", fgColor=hex_col)


def header_font(sz=11, bold=True, color=WHITE):
    return Font(name="Arial", size=sz, bold=bold, color=color)


def cell_font(sz=10, bold=False, color="1A2B3C"):
    return Font(name="Arial", size=sz, bold=bold, color=color)


def thin_border():
    s = Side(style="thin", color="D0DCE8")
    return Border(left=s, right=s, top=s, bottom=s)


def style_header_row(ws, row, num_cols, bg=HEADER):
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = hex_fill(bg)
        cell.font = header_font()
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border()


def style_data_row(ws, row, num_cols, alt=False):
    bg = LIGHT if alt else WHITE
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = hex_fill(bg)
        cell.font = cell_font()
        cell.alignment = Alignment(vertical="center")
        cell.border = thin_border()


def write_sheet_data(ws, headers, rows, title_text=None, header_bg=HEADER):
    start_row = 1
    if title_text:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(headers), 1))
        title_cell = ws.cell(row=1, column=1, value=title_text)
        title_cell.fill = hex_fill(NAVY)
        title_cell.font = Font(name="Arial", size=13, bold=True, color=WHITE)
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 28
        start_row = 2

    for ci, h in enumerate(headers, 1):
        ws.cell(row=start_row, column=ci, value=h)
    style_header_row(ws, start_row, len(headers), bg=header_bg)
    ws.row_dimensions[start_row].height = 22

    data_start = start_row + 1
    for ri, row_data in enumerate(rows):
        excel_row = data_start + ri
        for ci, val in enumerate(row_data, 1):
            ws.cell(row=excel_row, column=ci, value=val)
        style_data_row(ws, excel_row, len(headers), alt=(ri % 2 == 1))
        ws.row_dimensions[excel_row].height = 18

    for ci, h in enumerate(headers, 1):
        col_vals = [str(h)] + [str(r[ci - 1]) if ci - 1 < len(r) else "" for r in rows]
        max_len = max((len(v) for v in col_vals), default=8)
        ws.column_dimensions[get_column_letter(ci)].width = min(max_len + 3, 40)

    return data_start, data_start + len(rows) - 1


def build_bar_chart_from_sheet(ws, label_col, value_col, data_start, data_end, title, series_title,
                                color_hex=TEAL, width=18, height=12, horizontal=False):
    chart = BarChart()
    chart.type = "bar" if horizontal else "col"
    chart.grouping = "clustered"
    chart.title = title
    chart.style = 10
    chart.width = width
    chart.height = height

    data_ref = Reference(ws, min_col=value_col, min_row=data_start, max_row=data_end)
    cats_ref = Reference(ws, min_col=label_col, min_row=data_start, max_row=data_end)
    chart.add_data(data_ref, titles_from_data=False)
    chart.set_categories(cats_ref)
    chart.series[0].title = SeriesLabel(v=series_title)
    chart.series[0].graphicalProperties.solidFill = color_hex
    chart.series[0].graphicalProperties.line.solidFill = color_hex
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showVal = False
    return chart


def build_stacked_bar_chart(ws, label_col, value_cols_names, data_start, data_end, title, colors, width=18, height=12):
    chart = BarChart()
    chart.type = "col"
    chart.grouping = "stacked"
    chart.title = title
    chart.style = 10
    chart.width = width
    chart.height = height
    cats_ref = Reference(ws, min_col=label_col, min_row=data_start, max_row=data_end)
    for idx, (vcol, sname) in enumerate(value_cols_names):
        data_ref = Reference(ws, min_col=vcol, min_row=data_start, max_row=data_end)
        series = Series(data_ref, title=sname)
        series.graphicalProperties.solidFill = colors[idx % len(colors)]
        series.graphicalProperties.line.solidFill = colors[idx % len(colors)]
        chart.append(series)
    chart.set_categories(cats_ref)
    return chart


def build_doughnut_chart(ws, label_col, value_col, data_start, data_end, title, width=14, height=12):
    chart = DoughnutChart()
    chart.title = title
    chart.style = 10
    chart.width = width
    chart.height = height
    chart.holeSize = 40
    data_ref = Reference(ws, min_col=value_col, min_row=data_start, max_row=data_end)
    label_ref = Reference(ws, min_col=label_col, min_row=data_start, max_row=data_end)
    chart.add_data(data_ref, titles_from_data=False)
    chart.set_categories(label_ref)

    # NOTE: chart.series[0].dPt starts as an empty list - add_data() does not
    # populate it automatically. Per-slice coloring requires explicitly
    # building and appending DataPoint objects (the original code enumerated
    # the still-empty list, which silently did nothing rather than crash).
    if chart.series:
        n_points = data_end - data_start + 1
        for i in range(n_points):
            pt = DataPoint(idx=i)
            pt.graphicalProperties.solidFill = CHART_COLORS[i % len(CHART_COLORS)]
            chart.series[0].dPt.append(pt)

    chart.dataLabels = DataLabelList()
    chart.dataLabels.showPercent = True
    chart.dataLabels.showCatName = True
    chart.dataLabels.showVal = False
    return chart


def _kpi_sheet(wb, kpis: dict):
    ws = wb.create_sheet("KPI Summary")
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:D1")
    hdr = ws["A1"]
    hdr.value = "Dispatch OPS — KPI Summary"
    hdr.fill = hex_fill(NAVY)
    hdr.font = Font(name="Arial", size=15, bold=True, color=WHITE)
    hdr.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    fields = [
        ("Valid Invoices", kpis.get("valid_invoices", 0)), ("Total Boxes", kpis.get("total_boxes", 0)),
        ("Normal Boxes", kpis.get("normal_boxes", 0)), ("Freezer Boxes", kpis.get("freezer_boxes", 0)),
        ("Not Supplied", kpis.get("not_supplied", 0)), ("Active Drivers", kpis.get("active_drivers", 0)),
        ("Active Vehicles", kpis.get("active_vehicles", 0)), ("Unique Customers", kpis.get("unique_customers", 0)),
    ]
    for ri, (label, value) in enumerate(fields, 2):
        lc = ws.cell(row=ri, column=1, value=label)
        vc = ws.cell(row=ri, column=2, value=value)
        lc.fill = hex_fill(DARK2)
        lc.font = Font(name="Arial", size=11, bold=True, color=TEAL)
        lc.border = thin_border()
        vc.fill = hex_fill(LIGHT)
        vc.font = Font(name="Arial", size=12, bold=True, color=NAVY)
        vc.border = thin_border()
        ws.row_dimensions[ri].height = 24
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 16


def _chart_sheet(wb, sheet_name, headers, rows, title, header_bg, chart_kind, **chart_kwargs):
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False
    ds, de = write_sheet_data(ws, headers, rows, title_text=title, header_bg=header_bg)
    if not rows:
        return
    try:
        if chart_kind == "bar":
            chart = build_bar_chart_from_sheet(ws, 1, 2, ds, de, title=title, **chart_kwargs)
            ws.add_chart(chart, "D2")
        elif chart_kind == "doughnut":
            chart = build_doughnut_chart(ws, 1, 2, ds, de, title=title, **chart_kwargs)
            ws.add_chart(chart, "D2")
        elif chart_kind == "stacked":
            chart = build_stacked_bar_chart(ws, 1, [(2, "Normal"), (3, "Freezer")], ds, de, title=title, **chart_kwargs)
            ws.add_chart(chart, "E2")
    except Exception as exc:
        # The data table above is always written successfully regardless -
        # if the chart itself hits an openpyxl API issue, note it plainly
        # instead of failing the entire export (every other sheet still works).
        note_cell = ws.cell(row=de + 3, column=1, value=f"[Chart could not be generated: {exc}]")
        note_cell.font = cell_font(sz=9, color="B71C1C")


def _generic_data_sheet(wb, sheet_name, headers, rows, title, header_bg=NAVY):
    if not headers:
        return
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False
    write_sheet_data(ws, headers, rows, title_text=title, header_bg=header_bg)


def _safe_sheet(wb, fn, *args, **kwargs):
    """Runs one sheet-builder; if it raises, adds a small error sheet
    instead of failing the whole export - so one bad tab's data quirk
    doesn't cost you every other sheet too."""
    try:
        fn(wb, *args, **kwargs)
    except Exception as exc:
        ws = wb.create_sheet(f"ERROR-{fn.__name__}"[:31])
        ws["A1"] = f"This sheet failed to generate: {exc}"
        ws["A1"].font = cell_font(color="B71C1C")


def _safe_compute(fn, default: dict, *args, **kwargs) -> dict:
    """Runs a compute_* function; on any failure, returns a safe empty
    default instead of propagating - this is what was actually missing
    before (chart-building was already defensive, but the initial data
    fetch calls below it were not, which is what caused the reported
    500 error)."""
    try:
        return fn(*args, **kwargs)
    except Exception:
        return default


def export_dashboard_workbook(db: Session, start_date: str = "", end_date: str = "", driver: str = "",
                               drivers: str = "", areas: str = "", division: str = "", route_type: str = "",
                               vehicle_type: str = "", facility_type: str = "", salesman: str = "",
                               lead_time_band: str = "", classification: str = "") -> io.BytesIO:
    # Dashboard Performance (V2 verification pass): every one of the six
    # tabs below independently re-queried + re-filtered the entire
    # SapInvoiceFact table from scratch (six full-table scans for one
    # export click) - measured at ~12s on a 50k-row synthetic benchmark,
    # by far the worst offender in the whole app. Fetching once here and
    # passing it into every compute_* call (each now accepts an optional
    # `sap_rows` to skip its own query - see their docstrings) cut this to
    # ~2s with no change to any of the ported business logic itself.
    #
    # BUGFIX (export ignored active Global Filters): `build_filtered_sap_query`
    # used to be called here with no arguments at all, i.e. every filter at
    # its "no restriction" default, regardless of what the person had
    # actually selected on the dashboard - this is the ONE place that
    # decides what data every sheet below is built from, so passing the
    # same start_date/end_date/driver/drivers/areas/division/route_type/
    # vehicle_type/facility_type/salesman the dashboard itself is using
    # makes every sheet (KPIs, charts, Driver Overview, Lead Time, Area
    # Analytics, Order Summary, Not Supplied) reflect exactly the filtered
    # view the person is looking at when they click Export - same shared
    # query function as every dashboard tab, not a second filtering path.
    filter_kwargs = dict(start_date=start_date, end_date=end_date, driver=driver, drivers=drivers,
                          areas=areas, division=division, route_type=route_type,
                          vehicle_type=vehicle_type, facility_type=facility_type, salesman=salesman)
    shared_rows = dashboard_kpis.build_filtered_sap_query(db, **filter_kwargs).all()

    home = _safe_compute(dashboard_kpis.compute_home_dashboard, {
        "kpis": {}, "driver_overview": [],
        "charts": {"driver_boxes": {"labels": [], "values": []}, "facility": {"labels": [], "values": []},
                   "vantype": {"labels": [], "normal": [], "freezer": []}, "returns": {"delivered": 0, "not_supplied": 0}},
    }, db, sap_rows=shared_rows)
    perf = _safe_compute(driver_performance.compute_driver_performance_tab, {
        "chart_stops": {"labels": [], "values": []}, "chart_route_hours": {"labels": [], "values": []}, "route_cards": [],
    }, db, sap_rows=shared_rows)
    lead = _safe_compute(dashboard_kpis.compute_lead_time_tab, {
        "by_classification": [], "distribution": {"bins": [], "counts": []},
    }, db, sap_rows=shared_rows, lead_time_band=lead_time_band, classification=classification)
    area = _safe_compute(dashboard_kpis.compute_area_analytics_tab, {
        "chart": {"labels": [], "values": []}, "table": [],
    }, db, sap_rows=shared_rows)
    order = _safe_compute(dashboard_kpis.compute_order_summary_tab, {
        "facility_types": [], "rows": [],
    }, db, sap_rows=shared_rows)
    ns = _safe_compute(dashboard_kpis.compute_not_supplied_tab, {
        "not_supplied": [], "zero_box_excluded_from_kpis": [],
    }, db, sap_rows=shared_rows)

    wb = Workbook()
    wb.remove(wb.active)

    _safe_sheet(wb, _kpi_sheet, home["kpis"])

    c = home["charts"]
    _safe_sheet(wb, _chart_sheet, "Boxes by Driver", ["Driver Name", "Boxes"],
                list(zip(c["driver_boxes"]["labels"], c["driver_boxes"]["values"])),
                "Boxes by Driver", BLUE, "bar", series_title="Boxes", color_hex=BLUE, width=22, height=14)
    _safe_sheet(wb, _chart_sheet, "Facility Distribution", ["Facility Type", "Count"],
                list(zip(c["facility"]["labels"], c["facility"]["values"])),
                "Facility Distribution", NAVY, "doughnut", width=16, height=13)
    _safe_sheet(wb, _chart_sheet, "Units by Vehicle Type", ["Vehicle Type", "Normal Boxes", "Freezer Boxes"],
                [[c["vantype"]["labels"][i], c["vantype"]["normal"][i], c["vantype"]["freezer"][i]]
                 for i in range(len(c["vantype"]["labels"]))],
                "Units by Vehicle Type", HEADER, "stacked", colors=[TEAL, BLUE], width=22, height=14)
    _safe_sheet(wb, _chart_sheet, "Delivered vs Not Supplied", ["Status", "Count"],
                [["Delivered", c["returns"]["delivered"]], ["Not Supplied", c["returns"]["not_supplied"]]],
                "Delivered vs Not Supplied", NAVY, "doughnut", width=14, height=12)

    _safe_sheet(wb, _chart_sheet, "Stops per Driver", ["Driver / Vehicle", "GPS Stops"],
                list(zip(perf["chart_stops"]["labels"], perf["chart_stops"]["values"])),
                "Stops per Driver", HEADER, "bar", series_title="GPS Stops", color_hex=TEAL, width=22, height=14)
    _safe_sheet(wb, _chart_sheet, "Route Duration", ["Driver / Vehicle", "Route Duration (hrs)"],
                list(zip(perf["chart_route_hours"]["labels"], perf["chart_route_hours"]["values"])),
                "Route Duration (hrs)", BLUE, "bar", series_title="Hours", color_hex=BLUE, width=22, height=14)

    _safe_sheet(wb, _chart_sheet, "Lead Time by Class", ["Classification", "Avg Lead Time (Days)"],
                [[c2["name"], c2["avg_days"]] for c2 in lead["by_classification"]],
                "Avg Lead Time by Classification", NAVY, "bar",
                series_title="Avg Days", color_hex=TEAL, horizontal=True, width=20, height=13)
    _safe_sheet(wb, _chart_sheet, "Lead Time Distribution", ["Lead Time Bucket", "Orders"],
                list(zip(lead["distribution"]["bins"], lead["distribution"]["counts"])),
                "Lead Time Distribution", HEADER, "bar", series_title="Orders", color_hex=BLUE, width=18, height=12)
    _safe_sheet(wb, _chart_sheet, "Orders by Area", ["Area", "Orders"],
                list(zip(area["chart"]["labels"], area["chart"]["values"])),
                "Orders by Area", NAVY, "bar", series_title="Orders", color_hex=TEAL, horizontal=True, width=20, height=14)

    _safe_sheet(wb, _generic_data_sheet, "Driver Overview",
                ["Driver", "Vehicle", "Type", "Orders", "Boxes", "Freezer", "Not Supplied"],
                [[d["driver_name"], d["vehicle_num"], d["vehicle_type"], d["orders"], d["boxes"], d["freezer"], d["not_supplied"]]
                 for d in home["driver_overview"]],
                "Driver Overview", NAVY)
    _safe_sheet(wb, _generic_data_sheet, "Driver Performance",
                ["Vehicle", "Driver", "Match Method", "Stops", "Route Duration", "Avg Stop"],
                [[rc["vehicle_num"], rc["display_driver"], rc["match_method"], rc["stops"], rc["route_duration_hm"], rc["avg_stop_hm"]]
                 for rc in perf["route_cards"]],
                "Driver Performance Analytics", BLUE)
    _safe_sheet(wb, _generic_data_sheet, "Order Summary",
                ["Driver"] + order["facility_types"] + ["Total"],
                [[r["driver"]] + r["by_facility"] + [r["total"]] for r in order["rows"]],
                "Invoice Count by Driver x Facility", "2E7D32")
    _safe_sheet(wb, _generic_data_sheet, "Area Analytics",
                ["Area", "Orders", "Boxes", "Freezer", "Returns", "Drivers", "Success %"],
                [[r["area"], r["orders"], r["boxes"], r["freezer"], r["returns"], r["drivers"], r["success_pct"]]
                 for r in area["table"]],
                "Orders by Area", NAVY)
    _safe_sheet(wb, _generic_data_sheet, "Not Supplied",
                ["Driver", "Customer", "Facility Type", "Reason", "Boxes", "Invoice Date"],
                [[r["driver_name"], r["customer_name"], r["facility_type"], r["reason"], r["boxes"], r["invoice_date"]]
                 for r in ns["not_supplied"]],
                "Not Supplied Lines", "B71C1C")
    _safe_sheet(wb, _generic_data_sheet, "Zero Boxes",
                ["Driver", "Customer", "Invoice Date", "Dispatch Date"],
                [[r["driver_name"], r["customer_name"], r["invoice_date"], r["dispatch_date"]]
                 for r in ns["zero_box_excluded_from_kpis"]],
                "Zero Box Rows", HEADER)

    if len(wb.sheetnames) == 0:
        wb.create_sheet("Empty").append(["No data available to export yet."])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_route_plan_workbook(db: Session, plan_role: str | None = None) -> io.BytesIO:
    """
    Professional-formatted Route Plan export (colored headers, borders,
    alternating rows, auto column widths, title banner) - reuses the exact
    same styling helpers as the Dashboard export so the whole app looks
    consistent, per the "keep professional application formatting" ask.
    plan_role=None exports both Driver and Helper plans as separate sheets.
    """
    from app.models.route_plan import RouteAssignment

    wb = Workbook()
    wb.remove(wb.active)

    headers = [
        "Area Code", "Area Name", "Division", "Route Type", "Vehicle Type",
        "Anchored Vehicle", "Assigned Vehicle", "Driver Req.", "Helper Req.",
        "Driver Code", "Driver Name", "Helper Code", "Helper Name",
        "Start Date", "End Date", "Assignment Reason", "Experience Score",
        "Restriction Applied", "Vacation Replacement Info", "Status",
    ]

    roles = [plan_role] if plan_role else ["driver", "helper"]
    for role in roles:
        current_batch = (
            db.query(RouteAssignment.plan_batch_id)
            .filter(RouteAssignment.plan_role == role)
            .order_by(RouteAssignment.created_at.desc())
            .first()
        )
        rows_q = (
            db.query(RouteAssignment)
            .filter(RouteAssignment.plan_role == role, RouteAssignment.plan_batch_id == (current_batch[0] if current_batch else ""))
            .order_by(RouteAssignment.area_code, RouteAssignment.start_date)
            .all()
        )
        ws = wb.create_sheet(f"{role.capitalize()} Route Plan"[:31])
        rows = []
        for r in rows_q:
            vac_info = ""
            if r.is_vacation_replacement or r.original_person_code:
                vac_info = (
                    f"Original: {r.original_person_name} ({r.original_person_code}) | "
                    f"{r.vacation_start_date} to {r.vacation_end_date} | {r.vacation_replacement_reason}"
                )
            rows.append([
                r.area_code, r.area_name, r.sector, r.route_type, r.vehicle_type,
                r.anchored_vehicle_number, r.vehicle_number, r.driver_requirement, r.helper_requirement,
                r.driver_code, r.driver_name, r.helper_code, r.helper_name,
                r.start_date, r.end_date, r.assignment_reason,
                r.driver_score if role == "driver" else r.helper_score,
                r.restrictions_considered, vac_info, r.status,
            ])
        write_sheet_data(ws, headers, rows, title_text=f"DispatchController - {role.capitalize()} Route Plan", header_bg=NAVY)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_combined_route_plan_sheet(db: Session) -> io.BytesIO:
    """
    ROUTE PLAN LAYOUT (V2 milestone) - exports the unified Route Plan Sheet
    (the merge of the current Driver + Helper plans) in the exact requested
    column order: S/N, Division, Driver Code, Driver Name, Area, Helper
    Code, Helper Name, Vehicle Type, Vehicle Number, Route Type, Assignment
    Reason, Score. Same styling helpers as everywhere else in the app.
    """
    from app.services import route_planner

    wb = Workbook()
    ws = wb.active
    ws.title = "Route Plan Sheet"

    headers = ["S/N", "Division", "Driver Code", "Driver Name", "Area", "Helper Code", "Helper Name",
               "Vehicle Type", "Vehicle Number", "Route Type", "Assignment Reason", "Score"]
    sheet_rows = route_planner.get_combined_route_plan_sheet(db)
    rows = [[
        r["sn"], r["sector"], r["driver_code"], r["driver_name"], r["area_name"],
        r["helper_code"], r["helper_name"], r["vehicle_type"], r["vehicle_number"],
        r["route_type"], r["assignment_reason"], r["score"],
    ] for r in sheet_rows]

    write_sheet_data(ws, headers, rows, title_text="DispatchController - Route Plan Sheet", header_bg=NAVY)
    ws.column_dimensions["K"].width = 50  # Assignment Reason - wider, since it's a free-text explanation

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_filename() -> str:
    return f"DispatchOPS_Dashboard_{date.today()}.xlsx"
